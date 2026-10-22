#!/usr/bin/env python3
"""Реальные остатки ФФ из выгрузок их ВМС.

Загрузка выгрузки:  python3 ffreal.py load "ФФ №1" <файл.xlsx> [время ISO, по умолчанию сейчас]
Остаток = «Остаток свободный» из выгрузки минус заказы FBS по складам этого ФФ (оба кабинета),
созданные после времени выгрузки. Не считаются заказы, отменённые до того, как их собрали.
"""
import json, os, re, sys, time, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fetch import call, env, MSK

HERE = os.path.dirname(os.path.abspath(__file__))
SNAP = os.path.join(HERE, 'ff_snapshots.json')
MP = 'https://marketplace-api.wildberries.ru'
# склады ФФ по кабинетам: регулярка по имени склада
FF_WH = {'ФФ №1': r'^(фулфилмент №1|фф№1|фф №1)', 'ФФ2Б': r'^(?!dbs).*(ff2b|фф2б)', 'Флаконы': r'^флаконы', 'ONYX FF': r'^onyx', 'Граф': r'^граф', 'Самара один квадрат': r'^самара один квадрат', '3PL': r'^3pl'}


def base_article(a):
    """«Канистра 10,8л дубль» -> «Канистра 10,8л»: дубли берут остаток исходного артикула."""
    return re.sub(r'[\s_]*дубль\d*$', '', (a or '').strip(), flags=re.I)
CABS = [('erokl', 'TOKEN_2'), ('avanti', 'TOKEN_1')]
# Как называть строку выгрузки ВМС на сайте: {ФФ: {артикул в выгрузке: артикул продавца}}; одинаковые складываются
ALIASES = {'ФФ2Б': {'1260172441': 'Канистра 31л'}}
# ФФ, у которых в «Доступно» выгрузки ещё не учтены заказы в статусе «Новый»:
# при загрузке выгрузки такие заказы (созданные до её времени) фиксируются и тоже вычитаются
PRE_NEW_FF = {'ФФ2Б'}
# Артикулы из выгрузок ВМС, которые считаются комплектующими, а не товаром
COMPONENT_ARTICLES = {'ложка'}
CANCEL_WB = {'declined_by_client', 'canceled', 'canceled_by_client', 'canceled_by_missed_call'}


def load(ff, path, when=None):
    import openpyxl
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    h = [str(x or '').strip() for x in rows[0]]
    col = lambda *names: next((h.index(n) for n in names if n in h), None)
    ia = col('Артикул', 'Артикул (связь)', 'Артикул WB'); ib = col('Штрихкод базовый', 'Barcode', 'Штрихкод')
    ifree = col('Остаток свободный', 'Доступно'); itot = col('Остаток'); iname = col('Название', 'Наименование')
    iext = col('Штрихкоды дополнительные')
    if None in (ia, ifree):
        raise SystemExit(f'Не нашла колонки «Артикул» и «Остаток свободный»/«Доступно» в {h}')
    data, comps = {}, []
    for r in rows[1:]:
        if not r or str(r[0] or '').strip().lower() == 'итого': continue
        name = str((r[iname] if iname is not None else '') or '').strip()
        # комплектующие: упаковка без артикула и штучные короба (в т.ч. «из набора»)
        if not r[ia] or 'штучн' in name.lower() or str(r[ia]).strip().lower() in COMPONENT_ARTICLES:
            if name:
                comps.append({'name': name, 'article': str(r[ia] or '').strip(),
                              'total': int((r[itot] if itot is not None else 0) or 0),
                              'free': int(r[ifree] or 0),
                              'reserve': (int((r[itot] if itot is not None else 0) or 0) - int(r[ifree] or 0))})
            continue
        bcs = [str((r[ib] if ib is not None else '') or '').strip()] + ([b.strip() for b in str(r[iext] or '').split(',') if b.strip()] if iext is not None else [])
        a = str(r[ia]).strip()
        a = ALIASES.get(ff, {}).get(a, a)
        prev = data.get(a, {'free': 0, 'total': 0, 'barcodes': []})
        data[a] = {'free': prev['free'] + max(int(r[ifree] or 0), 0), 'total': prev['total'] + int((r[itot] if itot is not None else 0) or 0),
                   'barcodes': prev['barcodes'] + [b for b in bcs if b], 'title': (r[iname] if iname is not None else '') or ''}
    when = when or datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')
    snaps = json.load(open(SNAP)) if os.path.exists(SNAP) else {}
    snaps[ff] = {'time': when, 'file': os.path.basename(path), 'rows': data, 'components': comps}
    if ff in PRE_NEW_FF:
        snaps[ff]['pre'] = capture_new(ff, when)
        print(ff, 'зафиксировано новых заказов до выгрузки:', len(snaps[ff]['pre']))
    json.dump(snaps, open(SNAP, 'w'), ensure_ascii=False, indent=1)
    print(ff, 'загружено артикулов', len(data), 'время', when)


PITSUNOV = 'ИП Пицунов'


def load_pitsunov(path, when=None):
    """Отчёт «Остатки» ВМС по ИП Пицунов (.xls/.xlsx, шапка над таблицей, колонки Код/Артикул/Наименование/Доступно).
    Ключ — «Наименование» (это артикул продавца), штрихкод — «Код». Товар Аванти (ARTICLE_FILTER['avanti'])
    не берётся — он уже учтён по кабинету Аванти. Заказы FBS не вычитаются (своих складов в кабинетах нет)."""
    from fetch import ARTICLE_FILTER
    avanti = ARTICLE_FILTER.get('avanti', [])
    if path.lower().endswith('.xls'):
        import xlrd
        sh = xlrd.open_workbook(path).sheets()[0]
        rows = [sh.row_values(i) for i in range(sh.nrows)]
    else:
        import openpyxl
        rows = [list(r) for r in openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)]
    cell = lambda r, i: str(r[i] if i < len(r) and r[i] is not None else '').strip()
    hi = next(i for i, r in enumerate(rows) if 'Наименование' in [cell(r, j) for j in range(len(r))] and 'Доступно' in [cell(r, j) for j in range(len(r))])
    h = [cell(rows[hi], j) for j in range(len(rows[hi]))]
    ic, iname, ifree, itot = h.index('Код'), h.index('Наименование'), h.index('Доступно'), h.index('Остаток')
    whs = next((cell(r, j + 1) for r in rows[:hi] for j in range(len(r) - 1) if cell(r, j) == 'склад:'), '')
    data = {}
    for r in rows[hi + 1:]:
        a = cell(r, iname)
        if not a or a.lower().startswith('итого'): continue
        if any(m in a.lower() for m in avanti): continue
        num = lambda i: int(float(r[i] or 0)) if i < len(r) and str(r[i]).strip() != '' else 0
        bc = cell(r, ic).split('.')[0]
        prev = data.get(a, {'free': 0, 'total': 0, 'barcodes': []})
        data[a] = {'free': prev['free'] + max(num(ifree), 0), 'total': prev['total'] + num(itot),
                   'barcodes': prev['barcodes'] + ([bc] if bc else []), 'title': a}
    when = when or datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')
    snaps = json.load(open(SNAP)) if os.path.exists(SNAP) else {}
    snaps[PITSUNOV] = {'time': when, 'file': os.path.basename(path), 'rows': data, 'label': 'pitsunov', 'whs': whs}
    json.dump(snaps, open(SNAP, 'w'), ensure_ascii=False, indent=1)
    print(PITSUNOV, 'загружено артикулов', len(data), 'время', when)


def load_cabinet_rows(ff, path, cab, when, only=None):
    """Добавить в базу ФФ строки из отдельной выгрузки ВМС по другому кабинету (например, Аванти на ФФ2Б).
    У строк своё время выгрузки и свой кабинет: заказы считаются только этого кабинета и только после when.
    Заказы этого кабинета в статусе «Новый» на момент выгрузки тоже вычитаются (для ФФ из PRE_NEW_FF)."""
    import openpyxl
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    h = [str(x or '').strip() for x in rows[0]]
    ia, ib, ifree, itot, iname = h.index('Артикул'), h.index('Barcode'), h.index('Доступно'), h.index('Остаток'), h.index('Наименование')
    snaps = json.load(open(SNAP))
    sn = snaps[ff]
    # убрать прежние строки этого кабинета
    sn['rows'] = {a: r for a, r in sn['rows'].items() if r.get('cab') != cab}
    added = []
    for r in rows[1:]:
        if not r or not r[ia]: continue
        a = re.sub(r'_\s+', '_', str(r[ia]).strip())
        if only and not any(m in a.lower() for m in only): continue
        sn['rows'][a] = {'free': max(int(r[ifree] or 0), 0), 'total': int(r[itot] or 0),
                         'barcodes': [str(r[ib]).strip()] if r[ib] else [], 'title': r[iname] or '',
                         'cab': cab, 'since': when, 'file': os.path.basename(path)}
        added.append(a)
    if ff in PRE_NEW_FF:
        tok = env()[dict(CABS)[cab]]
        t = datetime.datetime.fromisoformat(when).astimezone(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%S')
        wl = call(tok, 'GET', f'{MP}/api/v3/warehouses')
        wids = {w['id'] for w in wl if re.search(FF_WH[ff], w['name'], re.I)}
        low = {a.lower() for a in added}
        sn['pre'] = [o for o in sn.get('pre', []) if o.get('cab') != cab]
        for o in call(tok, 'GET', f'{MP}/api/v3/orders/new').get('orders', []):
            if o['warehouseId'] in wids and o['createdAt'][:19] < t and (o.get('article') or '').lower() in low:
                sn['pre'].append({'cab': cab, 'id': o['id'], 'article': o.get('article'), 'skus': o.get('skus', []),
                                  'warehouseId': o['warehouseId'], 'createdAt': o['createdAt']})
    json.dump(snaps, open(SNAP, 'w'), ensure_ascii=False, indent=1)
    return added, [o for o in sn.get('pre', []) if o.get('cab') == cab]


def capture_new(ff, when):
    """Заказы в статусе «Новый» на складах ФФ, созданные до времени выгрузки (оба кабинета)."""
    e, out = env(), []
    t = datetime.datetime.fromisoformat(when).astimezone(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%S')
    for cab, key in CABS:
        tok = e.get(key)
        if not tok: continue
        wl = call(tok, 'GET', f'{MP}/api/v3/warehouses')
        wids = {w['id'] for w in wl if re.search(FF_WH[ff], w['name'], re.I)}
        for o in call(tok, 'GET', f'{MP}/api/v3/orders/new').get('orders', []):
            if o['warehouseId'] in wids and o['createdAt'][:19] < t:
                out.append({'cab': cab, 'id': o['id'], 'article': o.get('article'), 'skus': o.get('skus', []),
                            'warehouseId': o['warehouseId'], 'createdAt': o['createdAt']})
    return out


# --- Канистры Графа: остаток из 1С по семействам, только на исходном артикуле ---
CANISTER_FAMILIES = [  # (семейство, регулярка объёма в артикуле, исходный артикул)
    ('10л', r'10[,.]\d', 'Канистра 10,8л'),
    ('20л', r'2[01][,.]\d', 'Канистра 21,5'),
    ('30л', r'3[01]([,.]\d)?\s*л?', 'Канистра 31,5л'),
]
# С какого ФФ идёт поставка FBO с канистрами (по словам Лили). Остальные новые поставки с канистрами
# по умолчанию считаются поставками Графа.
from fetch import SUPPLY_FF as FBO_FROM
FAMILY_DEFAULT_FF = 'Граф'


def canister_units(article):
    """-> (семейство, штук канистр) или None. «Канистра 21,5л(2шт)» -> ('20л', 2)."""
    a = article or ''
    if 'анистр' not in a.lower():
        return None
    m = re.search(r'(\d+)\s*шт', a)
    n = int(m.group(1)) if m else 1
    for fam, rx, _ in CANISTER_FAMILIES:
        if re.search(r'анистр[аы]?\s*' + rx, a, re.I):
            return fam, n
    return None


def supply_units(sp):
    u = {}
    for g in sp['goods']:
        cu = canister_units(g['article'])
        if cu: u[cu[0]] = u.get(cu[0], 0) + cu[1] * g['qty']
    return u


def supply_ff(sid):
    return FBO_FROM.get(int(sid), FAMILY_DEFAULT_FF)


def load_family(ff, base, when, supplies, source='1С: ведомость по товарам на складах'):
    """Канистры по 1С. base: {семейство: штук}; supplies: текущие поставки FBO (берутся только поставки этого ФФ)."""
    snaps = json.load(open(SNAP)) if os.path.exists(SNAP) else {}
    fbo = {str(sp['id']): supply_units(sp) for sp in supplies
           if sp['cab'] == 'erokl' and sp['id'] and supply_ff(sp['id']) == ff and supply_units(sp)}
    snaps[ff] = {'kind': 'family', 'time': when, 'file': source, 'base': base, 'fbo': fbo,
                 'loaded': datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')}
    json.dump(snaps, open(SNAP, 'w'), ensure_ascii=False, indent=1)
    return snaps[ff]


def compute_family(ff, sn, e):
    """Остаток = 1С − поставки FBO с этого ФФ − заказы FBS по складам ФФ с момента sn['time']."""
    # поставки, которые теперь числятся за другим ФФ, убираем; новые поставки этого ФФ добавляем
    sn['fbo'] = {k: v for k, v in sn['fbo'].items() if supply_ff(k) == ff}
    dp = os.path.join(HERE, 'data.json')
    if os.path.exists(dp):
        for sp in json.load(open(dp))['supplies']:
            if sp['cab'] != 'erokl' or not sp['id'] or str(sp['id']) in sn['fbo'] or supply_ff(sp['id']) != ff: continue
            if sp['id'] not in FBO_FROM and (sp['created'] or '') < sn['loaded'][:19]: continue
            u = supply_units(sp)
            if u: sn['fbo'][str(sp['id'])] = u
    tok = e['TOKEN_2']
    wl = call(tok, 'GET', f'{MP}/api/v3/warehouses')
    wids = {w['id'] for w in wl if re.search(FF_WH[ff], w['name'], re.I)}
    orders, st = orders_since(tok, wids, int(datetime.datetime.fromisoformat(sn['time']).timestamp()))
    used = {f: 0 for f, _, _ in CANISTER_FAMILIES}
    for o in orders:
        s = st.get(o['id'], {})
        if s.get('supplierStatus') == 'cancel': continue
        if s.get('wbStatus') in CANCEL_WB and s.get('supplierStatus') in (None, 'new', 'confirm'): continue
        cu = canister_units(o.get('article'))
        if cu: used[cu[0]] += cu[1]
    fbo = {f: sum(v.get(f, 0) for v in sn['fbo'].values()) for f in used}
    rows = []
    for fam, _, art in CANISTER_FAMILIES:
        if fam not in sn['base']: continue
        b = sn['base'][fam]
        rows.append({'article': art, 'title': f'Канистры {fam}, все артикулы семейства', 'base': b, 'fbo': fbo[fam],
                     'orders': used[fam], 'raw': b - fbo[fam] - used[fam], 'now': max(b - fbo[fam] - used[fam], 0)})
    return rows


def orders_since(tok, wids, since_ts):
    out, nxt, now = [], 0, int(time.time())
    while True:
        r = call(tok, 'GET', f'{MP}/api/v3/orders?limit=1000&next={nxt}&dateFrom={since_ts}&dateTo={now}')
        out += [o for o in r.get('orders', []) if o['warehouseId'] in wids]
        if len(r.get('orders', [])) < 1000: break
        nxt = r['next']
    # новые заказы тоже (у них могут ещё не быть в /orders)
    new = call(tok, 'GET', f'{MP}/api/v3/orders/new').get('orders', [])
    seen = {o['id'] for o in out}
    for o in new:
        if o['warehouseId'] in wids and o['id'] not in seen and o['createdAt'] >= datetime.datetime.fromtimestamp(since_ts, datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%S'):
            out.append(o)
    st = {}
    ids = [o['id'] for o in out]
    for i in range(0, len(ids), 1000):
        for s in call(tok, 'POST', f'{MP}/api/v3/orders/status', {'orders': ids[i:i + 1000]}).get('orders', []):
            st[s['id']] = s
    return out, st


def compute():
    """-> {ff: {'time','file','rows':[{article,title,base,orders,now}], 'unmatched':{article:n}}}"""
    if not os.path.exists(SNAP): return {}
    snaps, e, res = json.load(open(SNAP)), env(), {}
    for ff, sn in snaps.items():
        if sn.get('kind') == 'family':
            rows = compute_family(ff, sn, e)
            res[ff] = {'time': sn['time'], 'file': sn['file'], 'rows': rows, 'unmatched': {}, 'exclusive': True,
                       'fbo': sn['fbo']}
            json.dump(snaps, open(SNAP, 'w'), ensure_ascii=False, indent=1)
            continue
        t = datetime.datetime.fromisoformat(sn['time'])
        since = int(t.timestamp())
        by_art = {a.lower(): a for a in sn['rows']}
        by_bc = {b: a for a, r in sn['rows'].items() for b in r['barcodes']}
        used, unmatched = {a: 0 for a in sn['rows']}, {}
        for cab, key in CABS:
            tok = e.get(key)
            if not tok: continue
            if ff not in FF_WH: continue  # своих складов в кабинетах нет (ИП Пицунов) — заказы не вычитаем
            wl = call(tok, 'GET', f'{MP}/api/v3/warehouses')
            wids = {w['id'] for w in wl if re.search(FF_WH[ff], w['name'], re.I)}
            if not wids: continue
            orders, st = orders_since(tok, wids, since)
            pre = [o for o in sn.get('pre', []) if o['cab'] == cab and o['id'] not in {x['id'] for x in orders}]
            if pre:
                ids = [o['id'] for o in pre]
                for i in range(0, len(ids), 1000):
                    for x in call(tok, 'POST', f'{MP}/api/v3/orders/status', {'orders': ids[i:i + 1000]}).get('orders', []):
                        st[x['id']] = x
                orders = orders + pre
            pre_ids = {o['id'] for o in pre}
            for o in orders:
                s = st.get(o['id'], {})
                if s.get('supplierStatus') == 'cancel': continue
                if s.get('wbStatus') in CANCEL_WB and s.get('supplierStatus') in (None, 'new', 'confirm'): continue
                art = (o.get('article') or '')
                a = by_art.get(art.lower()) or by_art.get(base_article(art).lower()) or next((by_bc[b] for b in o.get('skus', []) if b in by_bc), None)
                if a:
                    row = sn['rows'][a]
                    # строки из отдельной выгрузки кабинета (например, Аванти на ФФ2Б): свой кабинет и своё время
                    if row.get('cab') and row['cab'] != cab: continue
                    if row.get('since') and o['id'] not in pre_ids and o.get('createdAt', '')[:19] < row['since'][:19]: continue
                if a: used[a] += 1
                else: unmatched[o.get('article')] = unmatched.get(o.get('article'), 0) + 1
        rows = [{'article': a, 'title': r['title'], 'base': r['free'], 'orders': used[a], 'now': max(r['free'] - used[a], 0)}
                for a, r in sn['rows'].items()]
        res[ff] = {'time': sn['time'], 'file': sn['file'], 'rows': rows, 'unmatched': unmatched, 'pre': len(sn.get('pre', [])),
                   'components': sn.get('components', []),
                   'label': sn.get('label'), 'fbs_ok': sn.get('fbs_ok', False),
                   'whs': sn.get('whs')}
    return res


if __name__ == '__main__':
    if len(sys.argv) >= 4 and sys.argv[1] == 'load':
        load(sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else None)
    else:
        print(json.dumps(compute(), ensure_ascii=False, indent=1))
